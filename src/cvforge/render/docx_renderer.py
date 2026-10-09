"""DOCX version of the render view: single column, built-in styles, no tables or text boxes.

Parsers find structure through Word's own styles (Title, Heading 1, Heading 2, List Bullet),
and dates are right-aligned with a tab stop rather than a table.
"""

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from cvforge.render.segments import Segment

FONT = "Calibri"
MARGIN = Mm(16)
PAPER = {"a4": (Mm(210), Mm(297)), "us-letter": (Mm(215.9), Mm(279.4))}
_SEP = " | "
# style -> (size pt, bold, space before pt, space after pt)
_STYLES = {
    "Normal": (10.5, False, 0, 0),
    "Title": (22, True, 0, 2),
    "Heading 1": (11, True, 10, 3),
    "Heading 2": (10.5, True, 6, 0),
    "List Bullet": (10.5, False, 0, 0),
}


def _style(doc: Any, name: str, size: float, bold: bool, before: float, after: float) -> None:
    style = doc.styles[name]
    style.font.name, style.font.size, style.font.bold = FONT, Pt(size), bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), FONT)
    for theme in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{theme}"), None)  # a theme font would override ours
    fmt = style.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    borders = style.element.get_or_add_pPr().find(qn("w:pBdr"))
    if borders is not None:  # the default Title style draws a coloured rule
        borders.getparent().remove(borders)


def _run(paragraph: Paragraph, text: str, bold=False, italic=False, url: str | None = None) -> None:
    run = paragraph.add_run(text)
    run.bold, run.italic = bold or None, italic or None
    if url:  # python-docx has no writer for hyperlinks: wrap the run in w:hyperlink
        link = OxmlElement("w:hyperlink")
        relation = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link.set(qn("r:id"), relation)
        link.append(run._r)
        paragraph._p.append(link)


def _rich(paragraph: Paragraph, segments: list[Segment], italic: bool = False) -> None:
    for segment in segments:
        _run(paragraph, segment["text"], bool(segment.get("bold")), italic, segment.get("url"))


def _entry(doc: Any, block: dict[str, Any], right_edge: int) -> None:
    for line in block["lines"]:
        primary = line["style"] == "primary"
        paragraph = doc.add_paragraph(style="Heading 2" if primary else "Normal")
        _rich(paragraph, line["left"], italic=line["style"] == "secondary")
        if line["right"]:
            paragraph.paragraph_format.tab_stops.add_tab_stop(right_edge, WD_TAB_ALIGNMENT.RIGHT)
            run = paragraph.add_run("\t" + line["right"])
            run.bold = False
    for segments in block["paragraphs"]:
        _rich(doc.add_paragraph(), segments)
    for segments in block["bullets"]:
        _rich(doc.add_paragraph(style="List Bullet"), segments)


def render_docx(view: dict[str, Any], paper: str, out: Path) -> None:
    doc = Document()
    for name, spec in _STYLES.items():
        _style(doc, name, *spec)
    doc.styles["Heading 1"].font.all_caps = True  # display only: the text stays "Experience"
    section = doc.sections[0]
    section.page_width, section.page_height = PAPER[paper]
    section.left_margin = section.right_margin = MARGIN
    section.top_margin = section.bottom_margin = MARGIN
    right_edge = section.page_width - 2 * MARGIN
    doc.core_properties.title = f"{view['name']} Resume"
    doc.core_properties.author = view["name"]

    doc.add_paragraph(view["name"], style="Title")
    if view["headline"]:
        doc.add_paragraph(view["headline"])
    contact = doc.add_paragraph()  # in the body: header and footer text is often skipped
    for i, item in enumerate(view["contact"]):
        _run(contact, _SEP if i else "")
        _run(contact, item["text"], url=item["url"])

    for part in view["sections"]:
        doc.add_paragraph(part["heading"], style="Heading 1")
        for block in part["blocks"]:
            match block["kind"]:
                case "paragraph":
                    _rich(doc.add_paragraph(), block["segments"])
                case "labeled":
                    paragraph = doc.add_paragraph()
                    _run(paragraph, f"{block['label']}:", bold=True)
                    _run(paragraph, " ")
                    _rich(paragraph, block["segments"])
                case "bullets":
                    for item in block["items"]:
                        _rich(doc.add_paragraph(style="List Bullet"), item)
                case "inline":
                    paragraph = doc.add_paragraph()
                    for i, item in enumerate(block["items"]):
                        _run(paragraph, _SEP if i else "")
                        _rich(paragraph, item)
                case "entry":
                    _entry(doc, block, right_edge)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
