import json

import pytest
from docx import Document
from typer.testing import CliRunner

from cvforge.cli import app
from cvforge.config import Config, resource
from cvforge.export import to_jsonresume
from cvforge.models import Resume
from cvforge.render.docx_renderer import render_docx
from cvforge.render.text_renderer import render_text
from cvforge.render.view import build_view, flatten
from cvforge.textnorm import normalize
from cvforge.yaml_io import load_resume

runner = CliRunner()
EXAMPLE = resource("examples/resume.yaml")
NASTY = "C# & F# cost $5 < $9 > 1% of {x} \\ `code` a_b @home #42 * 2"


@pytest.fixture(scope="module")
def view():
    return build_view(load_resume(EXAMPLE), Config())


def in_order(expected: list[str], text: str) -> list[str]:
    """Items of `expected` that do not appear in `text` in order (normalized)."""
    hay, cursor, missing = normalize(text), 0, []
    for item in expected:
        found = hay.find(normalize(item), cursor)
        if found < 0:
            missing.append(item)
        else:
            cursor = found + len(normalize(item))
    return missing


def test_plain_text(view):
    text = render_text(view)
    lines = text.splitlines()
    assert lines[:2] == ["Jane Doe", "Senior Data Engineer"]
    assert lines[2].startswith("jane@example.com | +91 90000 00000 | Pune, India | linkedin.com")
    assert "EXPERIENCE" in lines and "Languages: Python, SQL, Bash" in lines
    assert "Example Corp | Pune, India" in lines
    assert "Senior Data Engineer | Jun 2023 – Present" in lines
    assert "**" not in text and "](" not in text  # no markup of any kind
    assert any(line.startswith("- Built AWS Glue and PySpark pipelines") for line in lines)
    assert in_order([t for _, t in flatten(view)], text) == []


def test_plain_text_shows_the_real_url_when_the_resume_shortens_it():
    resume = Resume.model_validate(
        {
            "basics": {
                "name": "Jane Doe",
                "links": [
                    {"url": "https://linkedin.com/in/jane-a1b2", "text": "linkedin.com/in/jane"}
                ],
            }
        }
    )
    view = build_view(resume, Config())
    assert "linkedin.com/in/jane-a1b2" in render_text(view)
    assert "[linkedin.com/in/jane](https://linkedin.com/in/jane-a1b2)" in render_text(view, True)


def test_markdown(view):
    text = render_text(view, markdown=True)
    lines = text.splitlines()
    assert lines[0] == "# Jane Doe"
    assert "## Experience" in lines and "## Skills" in lines
    assert "**Example Corp** | Pune, India  " in lines  # two spaces: a Markdown line break
    assert "*Senior Data Engineer* | Jun 2023 – Present" in lines
    assert "**Languages:** Python, SQL, Bash  " in lines
    assert any(line.startswith("- Built **AWS Glue** and **PySpark** pipelines") for line in lines)
    assert "**[lakehouse-lab](https://github.com/janedoe/lakehouse-lab)**  " in lines
    assert "[linkedin.com/in/janedoe](https://linkedin.com/in/janedoe)" in text


def test_docx_structure(view, tmp_path):
    out = tmp_path / "resume.docx"
    render_docx(view, "a4", out)
    doc = Document(str(out))
    styles = [p.style.name for p in doc.paragraphs]
    assert styles[0] == "Title" and doc.paragraphs[0].text == "Jane Doe"
    assert {"Heading 1", "Heading 2", "List Bullet", "Normal"} <= set(styles)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings[:3] == ["Summary", "Skills", "Experience"]  # standard words, not uppercased
    assert len(doc.tables) == 0 and len(doc.sections) == 1  # single column, no tables
    section = doc.sections[0]
    assert round(section.page_width.mm) == 210 and round(section.page_height.mm) == 297
    assert all(not p.text for p in section.header.paragraphs + section.footer.paragraphs)
    assert doc.styles["Normal"].font.name == "Calibri"
    assert 10 <= doc.styles["Normal"].font.size.pt <= 11
    assert doc.core_properties.author == "Jane Doe"


def test_docx_matches_the_pdf_content_and_order(view, tmp_path):
    out = tmp_path / "resume.docx"
    render_docx(view, "a4", out)
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert in_order([t for _, t in flatten(view)], text) == []
    # dates sit after a tab that a right tab stop aligns to the margin
    title = next(p for p in doc.paragraphs if p.text.startswith("Senior Data Engineer\t"))
    assert title.text.endswith("Jun 2023 – Present")
    assert [stop.alignment.name for stop in title.paragraph_format.tab_stops] == ["RIGHT"]
    links = {h.text: h.address for p in doc.paragraphs for h in p.hyperlinks}
    assert links["linkedin.com/in/janedoe"] == "https://linkedin.com/in/janedoe"
    assert links["lakehouse-lab"] == "https://github.com/janedoe/lakehouse-lab"
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "AWS Glue" in bold and "Languages:" in bold


def test_special_characters_survive_docx_txt_and_markdown(tmp_path):
    resume = Resume.model_validate(
        {
            "basics": {"name": "Jane O'Doe", "headline": NASTY},
            "summary": f"**{NASTY}** then {NASTY}",
            "experience": [
                {"company": "A < B > C & Co.", "title": "@Lead_Dev", "bullets": [NASTY]}
            ],
        }
    )
    view = build_view(resume, Config())
    out = tmp_path / "resume.docx"
    render_docx(view, "us-letter", out)
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert text.count(NASTY) == 4 and "A < B > C & Co." in text and "@Lead_Dev" in text
    assert render_text(view).count(NASTY) == 4
    assert render_text(view, markdown=True).count(NASTY) == 4


def test_render_writes_every_configured_format(tmp_path):
    result = runner.invoke(app, ["render", str(EXAMPLE), "--out", str(tmp_path)])
    assert result.exit_code == 0, result.output
    produced = {path.suffix for path in tmp_path.glob("Jane_Doe_Resume.*")}
    assert produced == {".pdf", ".docx", ".txt", ".md"}
    assert (tmp_path / "Jane_Doe_Resume.txt").read_text(encoding="utf-8").startswith("Jane Doe\n")


def test_formats_option(tmp_path):
    args = ["render", str(EXAMPLE), "--out", str(tmp_path)]
    result = runner.invoke(app, [*args, "--formats", "txt, md"])
    assert result.exit_code == 0 and "ATS check was skipped" in result.output
    assert {path.suffix for path in tmp_path.glob("Jane_Doe_Resume.*")} == {".txt", ".md"}
    assert runner.invoke(app, [*args, "--formats", "pdf,odt"]).exit_code == 2


def test_export_jsonresume():
    data = to_jsonresume(load_resume(EXAMPLE))
    assert (
        data["basics"]["name"] == "Jane Doe" and data["basics"]["label"] == "Senior Data Engineer"
    )
    assert data["basics"]["location"] == {"address": "Pune, India"}
    assert data["basics"]["profiles"][0] == {
        "network": "LinkedIn",
        "url": "https://linkedin.com/in/janedoe",
    }
    job = data["work"][0]
    assert (job["name"], job["position"], job["startDate"]) == (
        "Example Corp",
        "Senior Data Engineer",
        "2023-06",
    )
    assert "endDate" not in job  # a current role has no end date in JSON Resume
    assert job["highlights"][0].startswith("Built AWS Glue and PySpark")  # ** removed
    assert data["skills"][0] == {"name": "Languages", "keywords": ["Python", "SQL", "Bash"]}
    assert data["certificates"][0]["issuer"] == "Snowflake"
    assert data["meta"]["cvforge"]["extra_sections"][0]["title"] == "Volunteering"
    assert "awards" not in data  # empty sections are left out


def test_export_command(tmp_path):
    printed = runner.invoke(app, ["export", str(EXAMPLE), "--format", "jsonresume"])
    assert printed.exit_code == 0
    assert json.loads(printed.stdout)["basics"]["email"] == "jane@example.com"
    out = tmp_path / "resume.json"
    assert runner.invoke(app, ["export", str(EXAMPLE), "-o", str(out)]).exit_code == 0
    assert json.loads(out.read_text(encoding="utf-8"))["work"][0]["name"] == "Example Corp"
    assert runner.invoke(app, ["export", str(EXAMPLE), "--format", "xml"]).exit_code == 2
